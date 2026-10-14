import argparse
import sys

import openpyxl

from rag import config
from rag.documents import load_chunks
from rag.embeddings import embed_chunks_cached
from rag.vector_store import (
    create_payload_indexes,
    expand_context,
    get_qdrant_client,
    recreate_collection,
    search,
    upsert_chunks,
)

SEPARATOR = "-" * 80


def print_results(client, results, expand: bool = True) -> None:
    if not results:
        print("No results found.")
        return

    for r in results:
        p = r.payload
        print(SEPARATOR)
        print(f"Document : {p['document']}  (page {p['page']})")
        print(f"Score    : {r.score:.4f}")
        if expand:
            print(f"Context  : {expand_context(client, r)}")
        else:
            text = p["text"]
            print(f"Text     : {text[:500]}{'...' if len(text) > 500 else ''}")
    print(SEPARATOR)


def read_first_question(path: str, column: str, header_row: int, sheet: str | None) -> str | None:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.active
    for row in range(header_row + 1, ws.max_row + 1):
        value = ws[f"{column}{row}"].value
        if value and str(value).strip():
            return str(value).strip()
    return None


def cmd_ingest(args) -> None:
    chunks = load_chunks(args.pdf_dir)
    if not chunks:
        print("No PDFs found.")
        return

    client = get_qdrant_client()
    recreate_collection(client)
    embeddings = embed_chunks_cached(chunks)
    print("Uploading to Qdrant...")
    upsert_chunks(client, chunks, embeddings)
    print(f"Indexed {len(chunks)} chunks into '{config.COLLECTION_NAME}'.")


def cmd_create_indexes(args) -> None:
    create_payload_indexes(get_qdrant_client())
    print("Payload indexes created.")


def cmd_search(args) -> None:
    client = get_qdrant_client()
    results = search(client, args.question, top_k=args.top_k, folder=args.folder)
    print_results(client, results, expand=not args.no_expand)


def cmd_excel(args) -> None:
    question = read_first_question(args.path, args.column, args.header_row, args.sheet)
    if not question:
        sys.exit(f"No question found in column {args.column} of '{args.path}'.")

    print(f"Question: {question}\n")
    client = get_qdrant_client()
    print_results(client, search(client, question, top_k=args.top_k))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Semantic search over PDF documents with Azure OpenAI and Qdrant.")
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("ingest", help="Parse PDFs, embed them and (re)build the Qdrant collection.")
    p.add_argument("--pdf-dir", default=config.PDF_DIR)
    p.set_defaults(func=cmd_ingest)

    p = sub.add_parser("create-indexes", help="Create payload indexes on an existing collection.")
    p.set_defaults(func=cmd_create_indexes)

    p = sub.add_parser("search", help="Search the collection with a free-text question.")
    p.add_argument("question")
    p.add_argument("--top-k", type=int, default=5)
    p.add_argument("--folder", help="Restrict results to a sub-folder of the PDF directory.")
    p.add_argument("--no-expand", action="store_true", help="Show only the matched chunk, without neighbours.")
    p.set_defaults(func=cmd_search)

    p = sub.add_parser("excel", help="Search using the first question found in an Excel column.")
    p.add_argument("path")
    p.add_argument("--sheet")
    p.add_argument("--column", default="E")
    p.add_argument("--header-row", type=int, default=1)
    p.add_argument("--top-k", type=int, default=1)
    p.set_defaults(func=cmd_excel)

    return parser


if __name__ == "__main__":
    args = build_parser().parse_args()
    args.func(args)
